def export_docx(path, draft_document):
    try:
        from docx import Document
    except Exception:
        return None

    document = Document()
    for paragraph in draft_document.paragraphs:
        document.add_paragraph(paragraph.text)
    document.save(str(path))
    return path
